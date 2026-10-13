"""
Модуль для конвертации Excel файла в HTML страницы.
"""

import pandas as pd
import os
from openpyxl import load_workbook

# ФИКСИРОВАННЫЙ РАЗМЕР ШРИФТА - МЕНЯЙТЕ ЗДЕСЬ!
BASE_FONT_SIZE = "2.2em"      # Основной текст
HEADER_FONT_SIZE = "3.0em"    # Заголовки категорий

def debug_availability(df_display, sheet_name):
    """Полная диагностика столбца 'Наличие'"""
    print(f"\n     🔍 ДИАГНОСТИКА СТОЛБЦА 'НАЛИЧИЕ' (лист: {sheet_name}):")
    print(f"       Тип данных: {df_display['Наличие'].dtype}")
    print(f"       Всего значений: {len(df_display['Наличие'])}")
    
    # Выводим уникальные значения
    unique_vals = df_display['Наличие'].dropna().unique()
    print(f"       Уникальные значения ({len(unique_vals)}):")
    for val in unique_vals[:10]:  # первые 10
        print(f"         - '{val}' (тип: {type(val).__name__})")
    
    # Подробная диагностика по строкам
    print(f"\n       Подробно по строкам (первые 5):")
    availability_results = []
    
    for idx, value in df_display['Наличие'].items():
        # Определяем доступность
        is_available = True  # по умолчанию
        
        if pd.isna(value) or value is None:
            is_available = False
        elif isinstance(value, bool):
            is_available = value
        elif isinstance(value, (int, float)):
            # ВАЖНО: Pandas читает ИСТИНА/ЛОЖЬ как 1.0/0.0!
            is_available = bool(value)  # 0.0 -> False, 1.0 -> True
        elif isinstance(value, str):
            str_val = str(value).strip().lower()
            false_keywords = ['ложь', 'false', 'нет', '0', 'no', 'н']
            is_available = not any(keyword in str_val for keyword in false_keywords)
        
        availability_results.append(is_available)
        
        # Показываем только первые 5 для отладки
        if idx < 5:
            status = f"'{value}'"
            if pd.isna(value):
                status = "Пусто/NaN"
            print(f"       Строка {idx+1}: {status} -> доступно={is_available}")
    
    print(f"       ... и еще {len(df_display) - 5} строк")
    
    return pd.Series(availability_results, index=df_display.index)

def format_price(value):
    """Форматирование цены - убираем .0 и добавляем ₽"""
    if pd.isna(value):
        return ""
    
    # Если это число (float или int)
    if isinstance(value, (int, float)):
        # Преобразуем в целое, если это целое число
        if isinstance(value, float) and value.is_integer():
            return f"{int(value)}₽"
        else:
            # Для дробных чисел оставляем как есть
            return f"{value}₽"
    # Если это строка
    elif isinstance(value, str):
        # Убираем лишние пробелы
        cleaned = value.strip()
        # Убираем .0 в конце если есть
        if cleaned.endswith('.0'):
            cleaned = cleaned[:-2]
        # Добавляем ₽ если его нет
        if not cleaned.endswith('₽') and not cleaned.endswith('р') and not cleaned.endswith('руб'):
            cleaned += '₽'
        return cleaned
    else:
        return str(value)

def create_html_file(sheet_name, dataframe, background_image=None):
    """
    Создает HTML файл из DataFrame.
    """
    # Копируем DataFrame для обработки
    df_display = dataframe.copy()
    
    # Определяем столбцы для вывода
    ignore_columns = ['Наличие', 'Цвет текста', 'Цвет обводки']
    columns_to_show = [col for col in df_display.columns if col not in ignore_columns]
    
    if 'Название' not in df_display.columns:
        print(f"  Ошибка: В листе '{sheet_name}' нет столбца 'Название'!")
        return None
    
    # Обрабатываем столбец "Наличие" С ДИАГНОСТИКОЙ
    availability_series = None
    if 'Наличие' in df_display.columns:
        # Выполняем диагностику
        availability_series = debug_availability(df_display, sheet_name)
        
        # Статистика доступности
        print(f"     📊 Статистика доступности:")
        print(f"       Всего строк: {len(availability_series)}")
        print(f"       Доступно: {availability_series.sum()}")
        print(f"       Недоступно: {(~availability_series).sum()}")
    else:
        print(f"     ⚠ Столбец 'Наличие' не найден, все товары будут доступны")
        # Создаем серию из True
        availability_series = pd.Series([True] * len(df_display), index=df_display.index)
    
    # Получаем цвета
    text_colors = {}
    stroke_colors = {}
    
    if 'Цвет текста' in df_display.columns:
        for idx, val in df_display['Цвет текста'].items():
            if pd.isna(val) or str(val).strip() == '':
                text_colors[idx] = '#FFFFFF'
            else:
                color_val = str(val).strip()
                if not color_val.startswith('#'):
                    color_lower = color_val.lower()
                    if color_lower in ['белый', 'white']:
                        color_val = '#FFFFFF'
                    elif color_lower in ['черный', 'black']:
                        color_val = '#000000'
                    elif color_lower in ['красный', 'red']:
                        color_val = '#FF0000'
                    elif color_lower in ['зеленый', 'green']:
                        color_val = '#00FF00'
                    elif color_lower in ['синий', 'blue']:
                        color_val = '#0000FF'
                    elif color_lower in ['желтый', 'yellow']:
                        color_val = '#FFFF00'
                text_colors[idx] = color_val
    
    if 'Цвет обводки' in df_display.columns:
        for idx, val in df_display['Цвет обводки'].items():
            if pd.isna(val) or str(val).strip() == '':
                stroke_colors[idx] = '#000000'
            else:
                color_val = str(val).strip()
                if not color_val.startswith('#'):
                    color_lower = color_val.lower()
                    if color_lower in ['белый', 'white']:
                        color_val = '#FFFFFF'
                    elif color_lower in ['черный', 'black']:
                        color_val = '#000000'
                    elif color_lower in ['красный', 'red']:
                        color_val = '#FF0000'
                stroke_colors[idx] = color_val
    
    # CSS стиль
    css_style = f"""
    <style>
    /* ОСНОВНОЙ СТИЛЬ */
    * {{
        margin: 0;
        padding: 0;
        box-sizing: border-box;
    }}
    
    html, body {{
        width: 100vw;
        height: 100vh;
        overflow: hidden;
        margin: 0;
        padding: 0;
        font-family: 'Segoe UI', Arial, sans-serif;
        font-weight: 600;
    }}
    
    body {{
        padding: 20px 40px;
        display: flex;
        flex-direction: column;
    """
    
    if background_image and os.path.exists(background_image):
        css_style += f"""
        background: url('{background_image}') no-repeat center center fixed;
        background-size: cover;
        """
    else:
        css_style += """
        background: #ffffff;
        """
    
    css_style += f"""
    }}
    
    /* КОНТЕЙНЕР МЕНЮ */
    .menu-container {{
        flex: 1;
        display: flex;
        flex-direction: column;
        justify-content: flex-start;
        overflow: hidden;
    }}
    
    /* ЗАГОЛОВОК КАТЕГОРИИ - СЛЕВА */
    .category-header {{
        text-align: left;
        margin: 25px 0 10px 0;
        padding: 10px 0;
        flex-shrink: 0;
    }}
    
    .category-header h2 {{
        font-size: {HEADER_FONT_SIZE};
        margin: 0;
        font-weight: 700;
        display: block;
        padding: 5px 0;
        letter-spacing: 1px;
        text-transform: uppercase;
    }}
    
    /* СТРОКА БЛЮДА */
    .dish-row {{
        display: flex;
        justify-content: space-between;
        align-items: center;
        margin: 6px 0;
        padding: 4px 0;
        min-height: 1.8em;
        flex-shrink: 0;
    }}
    
    /* ЯЧЕЙКИ */
    .dish-cell {{
        padding: 0 8px;
        font-size: {BASE_FONT_SIZE};
        line-height: 1.2;
        flex: 1;
        white-space: nowrap;
        overflow: hidden;
        text-overflow: ellipsis;
    }}
    
    /* НАЗВАНИЕ БЛЮДА */
    .dish-cell:first-child {{
        text-align: left;
        font-weight: 600;
        flex: 3;
        padding-left: 0;
    }}
    
    /* ЦЕНА */
    .dish-cell:last-child {{
        text-align: right;
        font-weight: 700;
        min-width: 100px;
        flex: 0 0 auto;
        padding-right: 0;
    }}
    
    /* ОСТАЛЬНЫЕ СТОЛБЦЫ */
    .dish-cell:not(:first-child):not(:last-child) {{
        text-align: center;
        flex: 1;
        font-weight: 600;
    }}
    
    /* СТИЛЬ ДЛЯ ТЕКСТА С ОБВОДКОЙ */
    .text-with-outline {{
        display: inline-block;
    }}
    
    /* КЛАСС ДЛЯ ПРИМЕНЕНИЯ ОБВОДКИ */
    .text-stroke {{
        text-shadow: 
            1px 1px 2px var(--stroke, #000000),
            -1px -1px 2px var(--stroke, #000000),
            1px -1px 2px var(--stroke, #000000),
            -1px 1px 2px var(--stroke, #000000),
            0 0 4px var(--stroke, #000000) !important;
    }}
    
    /* НЕДОСТУПНЫЕ ПОЗИЦИИ - 50% ПРОЗРАЧНОСТИ */
    .dish-row.unavailable {{
        opacity: 0.5 !important;
        position: relative;
    }}
    
    .dish-row.unavailable .dish-cell {{
        color: #888888 !important;
        text-decoration: line-through !important;
        font-style: italic;
    }}
    
    .dish-row.unavailable::after {{
        content: '';
        position: absolute;
        top: 50%;
        left: 0;
        right: 0;
        height: 2px;
        background: linear-gradient(90deg, 
            transparent 5%, 
            rgba(255, 0, 0, 0.4) 50%, 
            transparent 95%);
        transform: translateY(-50%);
        pointer-events: none;
        opacity: 0.7;
    }}
    
    /* Стиль для заголовков с обводкой */
    .category-header h2.text-with-outline {{
        text-shadow: 
            2px 2px 3px var(--stroke, #000000),
            -2px -2px 3px var(--stroke, #000000),
            2px -2px 3px var(--stroke, #000000),
            -2px 2px 3px var(--stroke, #000000),
            0 0 6px var(--stroke, #000000) !important;
    }}
    </style>
    """
    
    # JavaScript
    javascript = """
    <script>
    // Запрет скролла
    document.addEventListener('wheel', function(e) {
        if (e.ctrlKey) return;
        e.preventDefault();
    }, { passive: false });
    
    // Скрываем курсор
    setTimeout(() => {
        document.body.style.cursor = 'none';
    }, 3000);
    
    // Отладка
    window.addEventListener('load', function() {
        console.log('=== ДЕБАГ ИНФОРМАЦИЯ ===');
        const unavailable = document.querySelectorAll('.dish-row.unavailable');
        console.log(`Недоступных позиций: ${unavailable.length}`);
        
        if (unavailable.length > 0) {
            unavailable.forEach((el, i) => {
                console.log(`❌ Недоступно ${i+1}: ${el.textContent.trim().substring(0, 50)}`);
            });
        } else {
            console.log('✅ Все позиции доступны');
        }
    });
    </script>
    """
    
    # HTML
    html_content = f"""<!DOCTYPE html>
<html lang="ru">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{sheet_name}</title>
    {css_style}
</head>
<body>
    <div class="menu-container">
"""
    
    # Читаем Excel для определения жирного шрифта
    excel_file = "BarMenu.xlsx"
    wb = load_workbook(excel_file, data_only=True)
    
    sheet_index = wb.sheetnames.index(sheet_name) if sheet_name in wb.sheetnames else 0
    ws = wb[wb.sheetnames[sheet_index]]
    
    start_row = 2
    
    # Обрабатываем каждую строку
    for i, row in df_display.iterrows():
        # Проверяем наличие
        is_available = True
        if availability_series is not None and i in availability_series:
            is_available = availability_series[i]
        
        # Берем цвета
        text_color = text_colors.get(i, '#FFFFFF')
        stroke_color = stroke_colors.get(i, '#000000')
        
        # Для отладки - показываем ТОЛЬКО недоступные
        if not is_available:
            dish_name = str(row.get('Название', '')).strip()
            print(f"     ⚠ Строка {i+1}: '{dish_name}' - НЕДОСТУПНО")
        
        # Проверяем жирный шрифт в ячейке "Название"
        cell_coord = ws.cell(row=start_row + i, column=1)  # Столбец A (Название)
        is_bold = cell_coord.font.bold if cell_coord.font else False
        
        if is_bold:
            # Заголовок категории
            category_name = str(row['Название']) if 'Название' in row else ""
            if category_name:
                html_content += f"""
        <div class="category-header">
            <h2 class="text-with-outline text-stroke" style="color: {text_color}; --stroke: {stroke_color};">{category_name}</h2>
        </div>
"""
        else:
            # Строка блюда
            row_class = "dish-row"
            if not is_available:
                row_class += " unavailable"
            
            html_content += f"""
        <div class="{row_class}">
"""
            
            # Добавляем все столбцы
            for j, col in enumerate(columns_to_show):
                cell_class = "dish-cell text-with-outline text-stroke"
                if j == 0:
                    cell_class += " truncate"
                elif j == len(columns_to_show) - 1:
                    cell_class += " right-align"
                
                cell_value = ""
                if col in row and not pd.isna(row[col]):
                    # ОСОБАЯ ОБРАБОТКА ДЛЯ КОЛОНКИ "ЦЕНА" и подобных
                    if col.lower() in ['цена', 'стоимость', 'price', 'cost']:
                        cell_value = format_price(row[col])
                    else:
                        cell_value = str(row[col])
                
                html_content += f"""
            <div class="{cell_class}" style="color: {text_color}; --stroke: {stroke_color};">{cell_value}</div>
"""
            
            html_content += """
        </div>
"""
    
    # Закрываем HTML
    html_content += f"""
    </div>
    {javascript}
</body>
</html>"""
    
    # Сохраняем файл
    filename = f"{sheet_name}.html"
    with open(filename, 'w', encoding='utf-8') as f:
        f.write(html_content)
    
    print(f"     💾 HTML файл сохранен: {filename}")
    return os.path.abspath(filename)

def convert_excel_to_html(excel_file="BarMenu.xlsx", background_image="Background.jpg"):
    """
    Основная функция конвертации.
    """
    print("=" * 70)
    print("🔄 КОНВЕРТАЦИЯ EXCEL В HTML")
    print("=" * 70)
    
    os.chdir(os.path.dirname(os.path.abspath(__file__)))
    
    if not os.path.exists(excel_file):
        print(f"❌ Ошибка: Файл {excel_file} не найден!")
        return []
    
    print(f"📊 Excel файл: {excel_file}")
    
    try:
        xls = pd.ExcelFile(excel_file)
        sheet_names = xls.sheet_names
        
        html_files = []
        
        for sheet_name in sheet_names:
            print(f"\n   📄 Лист: {sheet_name}")
            
            try:
                df = xls.parse(sheet_name, skiprows=0)
                
                if df.empty:
                    print(f"     ⚠ Пустой лист, пропускаем")
                    continue
                
                print(f"     📊 Строк: {len(df)}")
                
                # Показываем какие цвета найдены в Excel
                if 'Цвет текста' in df.columns:
                    colors = df['Цвет текста'].dropna().unique()
                    print(f"     🎨 Найдено цветов текста: {len(colors)}")
                
                if 'Цвет обводки' in df.columns:
                    strokes = df['Цвет обводки'].dropna().unique()
                    print(f"     🖌️ Найдено цветов обводки: {len(strokes)}")
                
                # Конвертируем
                html_path = create_html_file(sheet_name, df, background_image)
                
                if html_path:
                    html_files.append({
                        'name': f"{sheet_name}.html",
                        'path': html_path,
                        'sheet': sheet_name,
                        'url': "file:///" + html_path.replace('\\', '/')
                    })
                    print(f"     ✅ HTML создан")
                else:
                    print(f"     ❌ Ошибка создания HTML")
                    
            except Exception as e:
                print(f"     ❌ Ошибка: {e}")
                import traceback
                traceback.print_exc()
        
        print(f"\n📈 Всего создано HTML файлов: {len(html_files)}")
        
        if html_files:
            print("\n📋 Созданные файлы:")
            for file_info in html_files:
                print(f"   • {file_info['name']}")
        
        return html_files
        
    except Exception as e:
        print(f"\n❌ Ошибка при конвертации Excel: {e}")
        return []

if __name__ == "__main__":
    html_files = convert_excel_to_html()
    
    if html_files:
        print("\n" + "=" * 70)
        print("✅ КОНВЕРТАЦИЯ ЗАВЕРШЕНА")
        print("=" * 70)
        
        print(f"\n💡 Настройки шрифта:")
        print(f"   • Основной текст: {BASE_FONT_SIZE}")
        print(f"   • Заголовки: {HEADER_FONT_SIZE}")
        
        # УБРАЛИ ПРЕДЛОЖЕНИЕ ОТКРЫТЬ ФАЙЛЫ
    else:
        print("\n❌ Конвертация не удалась.")